import os
import re
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from openpyxl.styles import PatternFill
from comm_layer.diagnostic import Diagnostics
from comm_layer.uds import TEST_FAILED
from config.config import db_path, can_messages, crc_messages

FILL_PASS  = PatternFill(start_color='00B050', end_color='00B050', fill_type='solid')
FILL_FAIL  = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
FILL_ERROR = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')

EXCEL_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Template_TC_Design.xlsx')
SHEET_NAME = 'Diagnostics'

COL_TEST_CASE      = 1
COL_TEST_STEPS     = 5
COL_ACTUAL         = 7
COL_EXECUTION_TYPE = 8


def _check_dtc_inactive(bus, dtc_hex):
    dtcs = bus.check_dtc_by_status_mask()
    status = dtcs.get(int(dtc_hex, 16), 0) if dtcs else 0
    assert not (status & TEST_FAILED), f'DTC {dtc_hex} is active | Status: {status:02X}'


def _check_dtc_active(bus, dtc_hex):
    dtcs = bus.check_dtc_by_status_mask()
    status = dtcs.get(int(dtc_hex, 16), 0) if dtcs else 0
    assert status & TEST_FAILED, f'DTC {dtc_hex} is not active | Status: {status:02X}'


PATTERNS = [
    (r'Send request 19 02',                        lambda bus, m: bus.check_dtc_by_status_mask()),
    (r'Check DTC (0x[0-9A-Fa-f]+) is not active', lambda bus, m: _check_dtc_inactive(bus, m.group(1))),
    (r'Check DTC (0x[0-9A-Fa-f]+) is active',     lambda bus, m: _check_dtc_active(bus, m.group(1))),
    (r'Stop sending (\S+)',                         lambda bus, m: bus.stop_periodic_message(m.group(1))),
    (r'Wait (\d+\.?\d*)s',                         lambda bus, m: time.sleep(float(m.group(1)))),
    (r'Send periodic message (\S+)',                lambda bus, m: bus.send_periodic_message(m.group(1))),
    (r'Send periodic signal (\S+) = (\d+)',         lambda bus, m: bus.send_periodic_signal(m.group(1), int(m.group(2)))),
    (r'Populate results',                           lambda bus, m: None),
    (r'Verify the customer perception',             lambda bus, m: None),
]


def _execute_step(bus, step_text):
    for pattern, action in PATTERNS:
        m = re.search(pattern, step_text, re.IGNORECASE)
        if m:
            action(bus, m)
            return
    print(f'  [WARNING] Step not recognized: {step_text}')


def _run_test_case(bus, test_case_id, steps_text):
    print(f'\nRunning {test_case_id}...')
    steps = [line.strip() for line in steps_text.split('\n') if line.strip()]
    for step in steps:
        step_clean = re.sub(r'^\d+\.\s*', '', step)
        print(f'  -> {step_clean}')
        _execute_step(bus, step_clean)
    return 'PASS'


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]

    with Diagnostics(db_path) as bus:
        bus.build_crc_msgs(*crc_messages)
        bus.init_CAN_messages(can_messages)
        time.sleep(5)
        bus.clear_all_dtc()
        time.sleep(2)

        for row in ws.iter_rows(min_row=2):
            execution_type = row[COL_EXECUTION_TYPE].value
            test_case      = row[COL_TEST_CASE].value
            steps          = row[COL_TEST_STEPS].value

            if execution_type != 'Automated' or not test_case or not steps:
                continue

            result = 'ERROR: Did not execute'
            try:
                result = _run_test_case(bus, test_case, steps)
                print(f'  [PASS] {test_case}')
            except AssertionError as e:
                result = f'FAIL: {e}'
                print(f'  [FAIL] {test_case} | {e}')
            except Exception as e:
                result = f'ERROR: {e}'
                print(f'  [ERROR] {test_case} | {e}')

            cell = row[COL_ACTUAL]
            cell.value = result
            if result == 'PASS':
                cell.fill = FILL_PASS
            elif result.startswith('FAIL'):
                cell.fill = FILL_FAIL
            else:
                cell.fill = FILL_ERROR

    wb.save(EXCEL_PATH)
    print('\nResults written to Excel.')


if __name__ == '__main__':
    main()
